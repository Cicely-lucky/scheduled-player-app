#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Markdown → Word（简化版产品说明书）
"""
import os
import re
from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn

MD = os.path.join(os.path.dirname(__file__), "小云闹钟_产品使用说明书_简化版.md")
OUT = os.path.join(os.path.dirname(__file__), "小云闹钟_产品使用说明书_简化版.docx")
IMG_DIR = os.path.join(os.path.dirname(__file__), "screenshots")

doc = Document()

# 页面设置：A4
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)

# 默认字体
def set_default_font(run, font_name="宋体", size=10.5):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(size)

# 标题样式
styles = doc.styles
for lvl, font_name, size, bold in [
    ('Heading 1', '黑体', 16, True),
    ('Heading 2', '黑体', 13, True),
    ('Heading 3', '黑体', 11, True),
]:
    st = styles[lvl]
    st.font.name = font_name
    st._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    st.font.size = Pt(size)
    st.font.bold = bold
    st.paragraph_format.space_before = Pt(10)
    st.paragraph_format.space_after = Pt(6)

# 正文样式
body = styles['Normal']
body.font.name = '宋体'
body._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
body.font.size = Pt(10.5)
body.paragraph_format.line_spacing = 1.2
body.paragraph_format.space_after = Pt(4)

# 页脚：第 X 页 / 共 Y 页
footer = section.footer
footer_para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer_para.add_run()
footer_run.font.name = '宋体'
footer_run._element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
footer_run.font.size = Pt(9)
# 用字段实现动态页码
from docx.oxml import OxmlElement
fldChar1 = OxmlElement('w:fldChar')
fldChar1.set(qn('w:fldCharType'), 'begin')
instrText = OxmlElement('w:instrText')
instrText.text = 'PAGE'
fldChar2 = OxmlElement('w:fldChar')
fldChar2.set(qn('w:fldCharType'), 'end')
txt1 = OxmlElement('w:t'); txt1.text = '第 '
txt2 = OxmlElement('w:t'); txt2.text = ' 页 / 共 '
# 总页数字段
fldChar3 = OxmlElement('w:fldChar')
fldChar3.set(qn('w:fldCharType'), 'begin')
instrText2 = OxmlElement('w:instrText')
instrText2.text = 'NUMPAGES'
fldChar4 = OxmlElement('w:fldChar')
fldChar4.set(qn('w:fldCharType'), 'end')
txt3 = OxmlElement('w:t'); txt3.text = ' 页'
footer_run._r.append(txt1)
footer_run._r.append(fldChar1)
footer_run._r.append(instrText)
footer_run._r.append(fldChar2)
footer_run._r.append(txt2)
footer_run._r.append(fldChar3)
footer_run._r.append(instrText2)
footer_run._r.append(fldChar4)
footer_run._r.append(txt3)

def add_heading(text, level):
    p = doc.add_heading(level=level)
    run = p.add_run(text)
    set_default_font(run, '黑体', {1:16,2:13,3:11}[level])
    run.font.bold = True
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(6)
    return p

def add_para(text, bold=False, indent=False):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    p.paragraph_format.space_after = Pt(4)
    if indent:
        p.paragraph_format.left_indent = Cm(0.5)
    run = p.add_run(text)
    set_default_font(run, '宋体', 10.5, bold=bold)
    return p

def add_table(headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = h
        for p in hdr_cells[i].paragraphs:
            for r in p.runs:
                set_default_font(r, '宋体', 9.5)
                r.font.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v
            for p in cells[i].paragraphs:
                for r in p.runs:
                    set_default_font(r, '宋体', 9)
    doc.add_paragraph()

# 解析 Markdown
with open(MD, "r", encoding="utf-8") as fp:
    lines = fp.readlines()

i = 0
while i < len(lines):
    ln = lines[i].rstrip("\n")
    # 分页标记
    if ln.startswith("──── 第"):
        doc.add_page_break()
        i += 1
        continue
    # 标题
    m = re.match(r'^(#{1,3})\s+(.+)$', ln)
    if m:
        level = len(m.group(1))
        add_heading(m.group(2).strip(), level)
        i += 1
        continue
    # 图片 ![caption](path)
    m = re.match(r'^!\[(.*?)\]\((.*?)\)$', ln)
    if m:
        caption, path = m.group(1), m.group(2)
        abs_path = path if os.path.isabs(path) else os.path.normpath(os.path.join(os.path.dirname(MD), path))
        if os.path.exists(abs_path):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(abs_path, width=Cm(7.5))
            if caption:
                cp = doc.add_paragraph()
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cr = cp.add_run(caption)
                set_default_font(cr, '宋体', 9)
        i += 1
        # 下一行可能是 *图注*
        if i < len(lines):
            next_ln = lines[i].rstrip("\n")
            if next_ln.startswith("*") and next_ln.endswith("*"):
                cp = doc.add_paragraph()
                cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cr = cp.add_run(next_ln.strip("*"))
                set_default_font(cr, '宋体', 9)
                i += 1
        continue
    # 表格
    if ln.startswith("|"):
        # 收集连续表格行
        tbl_lines = []
        while i < len(lines) and lines[i].rstrip("\n").startswith("|"):
            tbl_lines.append(lines[i].rstrip("\n"))
            i += 1
        # 去掉分隔行
        content_rows = [r for r in tbl_lines if not re.match(r'^\|(\s*:?-+:?\s*\|)+\s*$', r.strip())]
        if content_rows:
            cells = [re.split(r'\s*\|\s*', r.strip("|")) for r in content_rows]
            headers = cells[0]
            rows = cells[1:]
            add_table(headers, rows)
        continue
    # 普通段落
    if ln.strip():
        # 处理粗体 **text**
        parts = re.split(r'(\*\*.*?\*\*)', ln)
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.2
        p.paragraph_format.space_after = Pt(4)
        # 判断是否是列表项
        list_match = re.match(r'^(\s*)([-*]|\d+\.)\s+', ln)
        if list_match:
            p.paragraph_format.left_indent = Cm(0.6)
        for part in parts:
            if part.startswith("**") and part.endswith("**"):
                run = p.add_run(part[2:-2])
                set_default_font(run, '宋体', 10.5)
                run.font.bold = True
            elif part:
                run = p.add_run(part)
                set_default_font(run, '宋体', 10.5)
    i += 1

doc.save(OUT)
print(f"[OK] Word 版已生成：{OUT}")
